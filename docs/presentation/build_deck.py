"""Build the Phase-1 project presentation in the ISE department template.

Geometry (bar height, footer positions, colours) was measured off the
department's own template PDF, so the output matches it rather than
approximating it. Cover and footer logo live in assets/.

Usage:  python docs/presentation/build_deck.py
"""

from copy import deepcopy

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
OUT = HERE / "Project31_Phase1_Presentation.pptx"

BAR_BLUE = RGBColor(0x44, 0x72, 0xC4)
BAR_EDGE = RGBColor(0x1F, 0x38, 0x64)
NAVY = RGBColor(0x00, 0x20, 0x60)
BLACK = RGBColor(0x00, 0x00, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

# Geometry measured from the department template (10 x 7.5 in page).
BAR_H = Inches(0.67)
BODY_L = Inches(0.55)
BODY_T = Inches(0.95)
BODY_W = Inches(8.9)
BODY_H = Inches(5.85)
LOGO_L, LOGO_T, LOGO_W = Inches(4.08), Inches(7.00), Inches(1.82)
FOOT_T = Inches(7.19)

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def kill_shadow(shape):
    """Drop the theme style reference and pin an empty effect list — without both,
    renderers fall back to the theme's drop shadow for autoshapes."""
    sp = shape._element
    for style in sp.findall(qn("p:style")):
        sp.remove(style)
    spPr = sp.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def flat_footer_box(box):
    """Zero the text insets so a footer line sits where it is placed."""
    tf = box.text_frame
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    tf.word_wrap = False


def header(slide, title):
    """Blue title bar across the top; empty title keeps the bar blank."""
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, BAR_H)  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = BAR_BLUE
    bar.line.color.rgb = BAR_EDGE
    bar.line.width = Pt(1)
    kill_shadow(bar)
    tf = bar.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = WHITE
    return bar


def footer(slide, page=None, dept=True):
    slide.shapes.add_picture(str(ASSETS / "footer_logo.png"), LOGO_L, LOGO_T, width=LOGO_W)
    if dept:
        box = slide.shapes.add_textbox(Inches(0.28), FOOT_T, Inches(3.2), Inches(0.3))
        flat_footer_box(box)
        p = box.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = "Department of ISE"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = NAVY
    if page is not None:
        box = slide.shapes.add_textbox(Inches(8.4), FOOT_T, Inches(1.28), Inches(0.3))
        flat_footer_box(box)
        p = box.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = f"{page:02d}"
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.name = FONT
        r.font.color.rgb = NAVY


def _bullet_char(p, char="•", indent_in=0.28):
    """python-pptx has no bullet API — set buChar and a hanging indent in XML."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(Inches(indent_in)).emu if False else int(Inches(indent_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": char})
    fnt = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    pPr.append(fnt)
    pPr.append(bu)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def body(slide, lines, size=18, space_after=12, line_spacing=1.0, top=None,
         height=None, bullet=True, indent_in=0.28):
    tb = slide.shapes.add_textbox(BODY_L, top or BODY_T, BODY_W, height or BODY_H)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if bullet:
            _bullet_char(p, indent_in=indent_in)
        else:
            _no_bullet(p)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = BLACK
    return tb


def numbered(slide, lines, size=18, space_after=12, top=None, height=None):
    tb = slide.shapes.add_textbox(BODY_L, top or BODY_T, BODY_W, height or BODY_H)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        p.line_spacing = 1.0
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.34))))
        pPr.set("indent", str(-int(Inches(0.34))))
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Calibri"}))
        # startAt only on the first item — repeating it restarts the list every line
        attrs = {"type": "arabicPeriod"}
        if i == 0:
            attrs["startAt"] = "1"
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), attrs))
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.name = FONT
        r.font.color.rgb = BLACK
    return tb


def fbox(slide, x, y, w, h, lines, size=11):
    """Flowchart node: white fill, thin black outline, black centred text."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = BLACK
    shp.line.width = Pt(1)
    kill_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.95
        p.space_after = Pt(0)
        _no_bullet(p)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size if i == 0 else size - 1)
        r.font.name = FONT
        r.font.color.rgb = BLACK
    return shp


def varrow(slide, xc, y1, y2):
    """Vertical connector with an arrowhead at the bottom end."""
    con = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                     Inches(xc), Inches(y1), Inches(xc), Inches(y2))
    con.line.color.rgb = BLACK
    con.line.width = Pt(1.25)
    ln = con.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"),
                             {"type": "triangle", "w": "med", "len": "med"}))
    return con


def centered(slide, lines, size, top, height=Inches(1.0), bold=False, color=BLACK,
             space_after=8, line_spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(0.5), top, Inches(9.0), height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        _no_bullet(p)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color
    return tb


# ---------------------------------------------------------------- slide 1: cover
s = add_slide()
s.shapes.add_picture(str(ASSETS / "cover.png"), 0, 0,
                     width=prs.slide_width, height=prs.slide_height)

# ------------------------------------------------------- slide 2: title / team
s = add_slide()
header(s, "")
centered(s, [
    "Cross-Domain Rumor Verification: Detecting Informed "
    "Trading Footprints via Sequential Decision-Making",
], 26, Inches(1.45), Inches(1.5), bold=True, line_spacing=1.15)
centered(s, ["Team Number: 31"], 20, Inches(3.30), Inches(0.4), bold=True)
centered(s, [
    "Sumukha Rao H  (NNM23IS189)",
    "Swati Prabhu  (NNM23IS192)",
    "Ujwal Hegde  (NNM23IS201)",
    "Vachan J Poojary  (NNM23IS205)",
], 18, Inches(3.95), Inches(1.8), space_after=5)
centered(s, ["Guide: Dr. Vaikunth Pai"], 20, Inches(5.95), Inches(0.4), bold=True)
footer(s, page=None, dept=False)

# ------------------------------------------------------------- slide 3: agenda
s = add_slide()
header(s, "AGENDA")
body(s, [
    "Introduction",
    "Literature Gap",
    "Motivation",
    "Problem Statement",
    "Objectives",
    "Methodology",
    "References",
], size=22, space_after=18, top=Inches(1.35))
footer(s, page=1)

# ------------------------------------------------------- slide 4: introduction
s = add_slide()
header(s, "INTRODUCTION")
body(s, [
    "Rumors appear on Reddit hours before official news — a few real, most noise.",

    "The stock often moves before the announcement becomes public.",

    "Verification is a timing problem: waiting adds evidence but spends the advantage.",

    "Our system re-decides every hour: WAIT, or COMMIT to TRUE or FALSE.",

    "Evidence spans two domains — Reddit activity and hourly price and volume.",
], size=20, space_after=24)
footer(s, page=2)

# ----------------------------------------------------- slide 5: literature gap
s = add_slide()
header(s, "LITERATURE GAP")
body(s, [
    "Existing rumor detection (Castillo 2011, Ma 2016, Zubiaga 2018) classifies a post "
    "once, at a fixed cut-off.",

    "It never chooses when to decide, and cannot abstain when evidence is thin.",

    "Social signals only — the market's reaction to the claim goes unused.",

    "Financial NLP predicts returns after news, not veracity before it.",

    "Earliness is a by-product of the chosen horizon, never an optimised objective.",

    "No dataset links a rumor's first post to the document that later confirms it.",
], size=20, space_after=22)
footer(s, page=3)

# ---------------------------------------------------------- slide 6: motivation
s = add_slide()
header(s, "MOTIVATION")
body(s, [
    "A verdict that arrives after the press release has no value.",

    "Abnormal price and volume before an announcement is a footprint of informed "
    "trading.",

    "The real trade-off is accuracy against hours saved.",

    "A wrong early verdict costs more than a slow one, so the agent may keep waiting.",

    "Free-tier data keeps the whole study reproducible on student hardware.",

    "Purpose is market-integrity monitoring — not trading advice.",
], size=20, space_after=22)
footer(s, page=4)

# --------------------------------------------------- slide 7: problem statement
s = add_slide()
header(s, "PROBLEM STATEMENT")
body(s, [
    "Input: a rumor event — ticker, claim, first-post time t₀ — with evidence "
    "arriving hourly.",

    "Each hour the agent picks one action: WAIT, COMMIT_TRUE or COMMIT_FALSE.",

    "Ground truth: t_official — the first SEC filing or whitelisted article within "
    "72 h.",

    "Goal: be correct, and be early — maximise Δ = t_official − t_commit.",

    "Constraint: features at step t use only data timestamped ≤ t₀ + t.",

    "Scope: binary veracity, US-listed tickers, 48-hour decision horizon.",
], size=20, space_after=22)
footer(s, page=5)

# ---------------------------------------------------------- slide 8: objectives
s = add_slide()
header(s, "OBJECTIVES")
numbered(s, [
    "Build a labelled dataset of rumor events with confirmation timestamps.",

    "Turn each event into an hourly state: text, social and market features.",

    "Model verification as a three-action MDP; train with PPO and DQN.",

    "Compare against fixed-horizon baselines and a zero-shot LLM policy.",

    "Evaluate accuracy/F1, calibration (Brier, ECE) and Time Delta Advantage.",

    "Demonstrate hour-by-hour replay of an event in a dashboard.",
], size=20, space_after=22)
footer(s, page=6)

# --------------------------------------------------------- slide 9: methodology
s = add_slide()
header(s, "METHODOLOGY")
COL_W, GAP = 2.83, 0.205
COL_X = [0.55, 0.55 + COL_W + GAP, 0.55 + 2 * (COL_W + GAP)]
COL_C = [x + COL_W / 2 for x in COL_X]
FULL_X, FULL_W = 0.55, 8.9
MID = FULL_X + FULL_W / 2

# three sources -> storage -> events -> labels -> state -> three models -> evaluation
fbox(s, COL_X[0], 0.92, COL_W, 0.62,
     ["Reddit posts", "Arctic Shift · 5 subreddits"])
fbox(s, COL_X[1], 0.92, COL_W, 0.62,
     ["Price & volume", "yfinance · hourly bars"])
fbox(s, COL_X[2], 0.92, COL_W, 0.62,
     ["News & filings", "GDELT · Finnhub · SEC EDGAR"])
for c in COL_C:
    varrow(s, c, 1.54, 1.86)

fbox(s, FULL_X, 1.86, FULL_W, 0.46,
     ["SQLite store — all timestamps in UTC   (106.9k posts · 3.7M bars)"])
varrow(s, MID, 2.32, 2.64)

fbox(s, FULL_X, 2.64, FULL_W, 0.46,
     ["Event construction — ticker matching → keyword filter → LLM triage → "
      "12 h clustering"])
varrow(s, MID, 3.10, 3.42)

fbox(s, FULL_X, 3.42, FULL_W, 0.46,
     ["Ground-truth labelling — TRUE / FALSE / UNVERIFIED, human reviewed   "
      "(1,170 rumor events)"])
varrow(s, MID, 3.88, 4.20)

fbox(s, FULL_X, 4.20, FULL_W, 0.46,
     ["Hourly state vector (418-d) — MiniLM + FinBERT · social · market"])
for c in COL_C:
    varrow(s, c, 4.66, 4.98)

fbox(s, COL_X[0], 4.98, COL_W, 0.66,
     ["RL agent (PPO / DQN)", "WAIT · COMMIT_TRUE · COMMIT_FALSE"])
fbox(s, COL_X[1], 4.98, COL_W, 0.66,
     ["Static baselines", "fixed horizons 6/12/24 h"])
fbox(s, COL_X[2], 4.98, COL_W, 0.66,
     ["LLM policy", "zero-shot, test split"])
for c in COL_C:
    varrow(s, c, 5.64, 5.96)

fbox(s, FULL_X, 5.96, FULL_W, 0.46,
     ["Evaluation — Accuracy / F1 · Brier · ECE · Time Delta Advantage Δ"])
footer(s, page=7)

# --------------------------------------------------------- slide 10: references
s = add_slide()
header(s, "REFERENCES")
numbered(s, [
    "Castillo, C., Mendoza, M., Poblete, B. (2011). Information Credibility on Twitter. "
    "WWW ’11, 675–684.",

    "Ma, J., Gao, W., Mitra, P., et al. (2016). Detecting Rumors from Microblogs with "
    "Recurrent Neural Networks. IJCAI 2016, 3818–3824.",

    "Ma, J., Gao, W., Wong, K.-F. (2017). Detect Rumors in Microblog Posts Using "
    "Propagation Structure via Kernel Learning. ACL 2017, 708–717.",

    "Zubiaga, A., Aker, A., Bontcheva, K., et al. (2018). Detection and Resolution of "
    "Rumours in Social Media: A Survey. ACM Computing Surveys, 51(2), 1–36.",

    "Zhou, K., Shu, C., Li, B., Lau, J. H. (2019). Early Rumour Detection. NAACL-HLT "
    "2019, 1614–1623.",

    "Vosoughi, S., Roy, D., Aral, S. (2018). The Spread of True and False News Online. "
    "Science, 359(6380), 1146–1151.",

    "Schulman, J., Wolski, F., Dhariwal, P., et al. (2017). Proximal Policy Optimization "
    "Algorithms. arXiv:1707.06347.",
], size=16, space_after=12)
footer(s, page=8)

# --------------------------------------------------------- slide 11: thank you
s = add_slide()
header(s, "")
centered(s, ["THANK YOU"], 40, Inches(3.1), Inches(0.9), bold=True)
footer(s, page=9)

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
